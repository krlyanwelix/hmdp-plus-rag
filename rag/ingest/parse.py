"""文档解析：任意格式 -> 纯文本 + 结构锚点。

依赖策略（缺依赖报错信息会直接告诉你装什么）：
    .txt/.md        零依赖
    .docx           python-docx
    .pdf            pypdf
    .html/.htm      beautifulsoup4

注意：解析质量决定整个 RAG 的天花板。解析完请人工抽查前几页输出，
如果人读着都乱，模型更读不懂。
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from rag.types import Document

SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".docx", ".pdf", ".html", ".htm"}

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$", re.MULTILINE)


def normalize(text: str) -> str:
    """统一换行、去掉行尾空白、压缩连续空行。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _doc_id(path: Path) -> str:
    rel = path.as_posix()
    return hashlib.blake2b(rel.encode("utf-8"), digest_size=8).hexdigest()


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _md_anchors(text: str) -> list[tuple[int, str]]:
    """提取 markdown 标题及其字符偏移。"""
    out: list[tuple[int, str]] = []
    for m in _MD_HEADING.finditer(text):
        out.append((m.start(), m.group(2).strip()))
    return out


def _read_docx(path: Path) -> tuple[str, list[tuple[int, str]]]:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise ImportError("解析 .docx 需要 python-docx，请执行：pip install python-docx") from exc

    d = docx.Document(str(path))
    parts: list[str] = []
    anchors: list[tuple[int, str]] = []
    cursor = 0
    for para in d.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading") or style in {"title", "subtitle"}:
            anchors.append((cursor, txt))
        parts.append(txt)
        cursor += len(txt) + 2
    return "\n\n".join(parts), anchors


def _read_pdf(path: Path) -> tuple[str, list[tuple[int, str]]]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise ImportError("解析 .pdf 需要 pypdf，请执行：pip install pypdf") from exc

    reader = PdfReader(str(path))
    parts: list[str] = []
    anchors: list[tuple[int, str]] = []
    cursor = 0
    for i, page in enumerate(reader.pages, start=1):
        page_text = normalize(page.extract_text() or "")
        if not page_text:
            continue
        anchors.append((cursor, f"第 {i} 页"))
        parts.append(page_text)
        cursor += len(page_text) + 2
    return "\n\n".join(parts), anchors


def _read_html(path: Path) -> tuple[str, list[tuple[int, str]]]:
    try:
        from bs4 import BeautifulSoup  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise ImportError("解析 .html 需要 beautifulsoup4，请执行：pip install beautifulsoup4") from exc

    soup = BeautifulSoup(_read_text(path), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()

    body = soup.body or soup
    parts: list[str] = []
    anchors: list[tuple[int, str]] = []
    cursor = 0
    for el in body.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "td"]):
        txt = normalize(el.get_text(" ", strip=True))
        if not txt:
            continue
        if el.name.startswith("h") and len(el.name) == 2:
            anchors.append((cursor, txt))
        parts.append(txt)
        cursor += len(txt) + 2
    return "\n\n".join(parts), anchors


def parse_file(path: Path, base: Path | None = None) -> Document:
    """解析单个文件。base 用于生成相对路径（引用溯源更可读）。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的格式 {suffix}：{path}")

    if suffix in {".md", ".markdown"}:
        text = normalize(_read_text(path))
        anchors = _md_anchors(text)  # 必须在 normalize 之后算偏移
    elif suffix == ".txt":
        text, anchors = normalize(_read_text(path)), []
    elif suffix == ".docx":
        text, anchors = _read_docx(path)
        text = normalize(text)
    elif suffix == ".pdf":
        text, anchors = _read_pdf(path)
        text = normalize(text)
    else:
        text, anchors = _read_html(path)
        text = normalize(text)

    try:
        source = path.relative_to(base).as_posix() if base else path.name
    except ValueError:
        source = path.name

    return Document(
        doc_id=_doc_id(path.resolve()),
        source=source,
        text=text,
        metadata={"path": str(path.resolve()), "suffix": suffix},
        anchors=anchors,
    )


def load_dir(directory: Path | str) -> list[Document]:
    """递归加载目录下所有支持的格式。"""
    directory = Path(directory)
    if not directory.exists():
        raise FileNotFoundError(f"原始文档目录不存在：{directory}")

    docs: list[Document] = []
    for p in sorted(directory.rglob("*")):
        if not p.is_file() or p.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if p.name.startswith("."):
            continue
        try:
            doc = parse_file(p, base=directory)
        except Exception as exc:  # 单文件失败不应中断整批索引
            print(f"  [跳过] {p.name}: {exc}")
            continue
        if doc.text:
            docs.append(doc)
    return docs
